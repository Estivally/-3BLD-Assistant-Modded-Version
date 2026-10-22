#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
alg818_qm_import.py
===================
把「浅梦」202602 版盲拧公式集（棱块 UF 分类 / 角块 UFR 分类）转换成
三盲助手 App 可直接读取的资产文件：

    只读  <recv>/浅梦's 棱块UF分类公式集 202602版.xlsx
    只读  <recv>/浅梦's 角块UFR分类公式集 202602版.xlsx
    只读  src/assets/alg818.txt
    只读  src/assets/assets/json/cornerPosToCode.json / edgePosToCode.json
    只写  src/assets/alg818_qm.txt

幂等：同一输入重复运行输出逐字节一致（无时间戳、稳定排序、LF、UTF-8 无 BOM）。

输出每行 9 个字段（Tab 分隔）：
    TYPE  KEY  CAT  HAND  COMM  NOTE  ALG  ALT  ALT_COMM

用法：
    python tools/alg818_qm_import.py            # 生成 + 自检
    python tools/alg818_qm_import.py --check    # 只自检，不写文件
    python tools/alg818_qm_import.py --quiet    # 只打印结论行

设计要点
--------
1. 列名不统一 → 一律按表头名取列；表头缺失的列按「备选 +1/+2/+3」位置推断。
2. 同一 case 会同时出现在「总表」与「分类页」→ 按编码去重，CAT 取并集。
3. 浅梦编码是 2 字母（缓冲省略），App 键是「缓冲字母 + 2 字母」。
   · 角块：浅梦缓冲 = J(UFR)，App 缓冲 = J(UFR) → f_corner = 恒等（投票已确认）。
   · 棱块：浅梦缓冲 = A(UF)，App 缓冲 = G(UR)。App 键里不允许出现 G/H，
     而浅梦编码里 G/H 是普通字母，所以必须 A/B ↔ G/H 对调、其余 20 字母恒等；
     这样既满足双射，又让学员已背的 2 字母词对保持原样。
4. 映射由「公式投票」推导（角块命中即采用）；棱块投票会给出非双射，脚本明确
   拒绝并说明原因，改用「缓冲对应 + 双射约束」，再用两道硬校验兜底：
   · 818 个 KEY 两两不同，且与 alg818.txt 的键集合完全相等；
   · 818 条公式逐个用魔方模拟求出 3-cycle，必须与「缓冲位 + 编码」完全吻合
     （这一步同时逐字母验证了「编码」页字母图与 App 的 posToCode 一致）。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
from collections import Counter, OrderedDict, defaultdict
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

import openpyxl

# --------------------------------------------------------------------------- #
# 路径
# --------------------------------------------------------------------------- #
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
ASSETS = os.path.join(ROOT, "src", "assets")
JSOND = os.path.join(ASSETS, "assets", "json")
ALG818 = os.path.join(ASSETS, "alg818.txt")
OUTFILE = os.path.join(ASSETS, "alg818_qm.txt")

RECV = os.environ.get("QM_RECV_DIR", r"D:\Users\Estiavl\Documents\Tencent Files\1103094647\FileRecv")
XLSX_EDGE = os.path.join(RECV, "浅梦's 棱块UF分类公式集 202602版.xlsx")
XLSX_CORNER = os.path.join(RECV, "浅梦's 角块UFR分类公式集 202602版.xlsx")

# 分类：严格采用任务书枚举值；sheet 名 → CAT 的归一表见 CAT_ALIAS
CORNER_CATS = ["Base", "DU", "UU", "R.U.", "R.D.", "R.F(B).R.", "Others", "BT"]
EDGE_CATS = ["U8", "E8", "S8", "M8", "E4", "S4", "M4", "Others"]
CAT_ALIAS = {"R.U.(R.)": "R.U.", "R.D.(R.)": "R.D."}
CAT_ORDER = {"C": CORNER_CATS, "E": EDGE_CATS}
MASTER_SHEETS = {"UF总表", "UFR总表"}
SKIP_SHEETS = {"说明", "编码"}
FALLBACK_CAT = "Others"

# 任务书给出的语义抽检期望（用于报告，不用于放宽校验）
CAT_HINT = {
    "U8": "U 层直达（两字母含 G 且相邻 / 不含 G 为 112、121 型）",
    "E8": "空穴 FR/FL，交换步骤在 E 层",
    "S8": "空穴 UR/UL，交换步骤在 S 层",
    "M8": "空穴 UF/UB，交换步骤在 M 层",
    "E4": "4 步公式，三块均挪到 E 层",
    "S4": "4 步公式，S 层方向同 E4",
    "M4": "4 步公式，底层字母打到对角线",
    "Others": "混合类",
    "Base": "基础 16 条（DD' 相间）",
    "DU": "下起手",
    "UU": "上起手",
    "R.U.": "R U' R' setup 类",
    "R.D.": "R D' R' setup 类",
    "R.F(B).R.": "R' F R setup 类",
    "BT": "肌肉记忆类",
}

# --------------------------------------------------------------------------- #
# 文本归一
# --------------------------------------------------------------------------- #
_QUOTES = str.maketrans({"\u2019": "'", "\u2018": "'", "\u2032": "'",
                         "\uff07": "'", "\u00b4": "'", "\u2035": "'"})
_WS = re.compile(r"\s+")


def norm_text(s) -> str:
    """换行/制表符/全角空格 → 单空格，统一撇号，去首尾空白。"""
    if s is None:
        return ""
    s = str(s).translate(_QUOTES)
    s = s.replace("\u3000", " ").replace("\t", " ").replace("\r", " ").replace("\n", " ")
    return _WS.sub(" ", s).strip()


def loose(s: str) -> str:
    """公式比对用：去空白。"""
    return norm_text(s).replace(" ", "")


def is_placeholder(s: str) -> bool:
    """blddb 的 'Not found.' 占位。"""
    t = s.strip().lower().rstrip(".")
    return t in ("not found", "notfound", "none", "n/a", "-", "—", "无")


# --------------------------------------------------------------------------- #
# 魔方模拟（坐标：x=R+ / y=F+ / z=U+）
# --------------------------------------------------------------------------- #
def _rz90(p):  return (-p[1], p[0], p[2])
def _rzm90(p): return (p[1], -p[0], p[2])
def _rx90(p):  return (p[0], -p[2], p[1])
def _rxm90(p): return (p[0], p[2], -p[1])
def _ry90(p):  return (p[2], p[1], -p[0])
def _rym90(p): return (-p[2], p[1], p[0])


NORMALS = {"U": (0, 0, 1), "D": (0, 0, -1), "R": (1, 0, 0),
           "L": (-1, 0, 0), "F": (0, 1, 0), "B": (0, -1, 0)}
NORMAL_INV = {v: k for k, v in NORMALS.items()}

# 面/切片 → (旋转函数, 层判定)
MOVES = {
    "U": (_rz90,  lambda p: p[2] == 1),
    "D": (_rzm90, lambda p: p[2] == -1),
    "R": (_rx90,  lambda p: p[0] == 1),
    "L": (_rxm90, lambda p: p[0] == -1),
    "F": (_ry90,  lambda p: p[1] == 1),
    "B": (_rym90, lambda p: p[1] == -1),
    "M": (_rxm90, lambda p: p[0] == 0),
    "E": (_rzm90, lambda p: p[2] == 0),
    "S": (_ry90,  lambda p: p[1] == 0),
}
# 宽层（u/d/r/l/f/b）= 对应面 + 中间层
WIDE = {
    "u": ("U", lambda p: p[2] >= 0), "d": ("D", lambda p: p[2] <= 0),
    "r": ("R", lambda p: p[0] >= 0), "l": ("L", lambda p: p[0] <= 0),
    "f": ("F", lambda p: p[1] >= 0), "b": ("B", lambda p: p[1] <= 0),
}
ROT = {"x": _rx90, "y": _ry90, "z": _rz90}

SLOTS: List[Tuple[Tuple[int, int, int], str]] = []
for _x in (-1, 0, 1):
    for _y in (-1, 0, 1):
        for _z in (-1, 0, 1):
            if abs(_x) + abs(_y) + abs(_z) in (2, 3):
                for _f, _n in NORMALS.items():
                    if _n[0] * _x + _n[1] * _y + _n[2] * _z == 1:
                        SLOTS.append(((_x, _y, _z), _f))

_TOKEN_RE = re.compile(r"^([UDLRFBMESudlrfbxyz])(w?)(2?)('?)$")


def parse_token(tok: str):
    m = _TOKEN_RE.match(tok)
    if not m:
        return None
    return m.group(1), m.group(2), m.group(3), m.group(4)


def apply_move(tok: str) -> Dict[Tuple, Tuple]:
    """返回 sticker → 目标 sticker 的映射。"""
    f, w, c, p = parse_token(tok)
    if f in ROT:
        rot, layer = ROT[f], (lambda q: True)
    elif f in WIDE and not w:
        rot, layer = MOVES[WIDE[f][0]][0], WIDE[f][1]
    else:
        rot, layer = MOVES[f]
    n = 2 if c == "2" else (3 if p == "'" else 1)
    out = {}
    for slot in SLOTS:
        pos, face = slot
        if layer(pos):
            np = pos
            for _ in range(n):
                np = rot(np)
            nv = NORMALS[face]
            for _ in range(n):
                nv = rot(nv)
            out[slot] = (np, NORMAL_INV[nv])
        else:
            out[slot] = slot
    return out


def alg_perm(alg: str) -> Dict[Tuple, Tuple]:
    perm = {s: s for s in SLOTS}
    for tok in alg.split():
        mv = apply_move(tok)
        perm = {s: mv[perm[s]] for s in perm}
    return perm


def cycles_of(perm) -> List[List[Tuple]]:
    seen, out = set(), []
    for s in SLOTS:
        if s in seen:
            continue
        cyc, cur = [], s
        while cur not in seen:
            seen.add(cur)
            cyc.append(cur)
            cur = perm[cur]
        if len(cyc) > 1:
            out.append(cyc)
    return out


def build_letter_slots(pos2code: Dict[str, str]) -> Dict[Tuple, str]:
    """posToCode 的 'UB':E 表示「UB 块的 U 面贴纸是 E」。"""
    m = {}
    for pos, code in pos2code.items():
        x = 1 if "R" in pos else (-1 if "L" in pos else 0)
        y = 1 if "F" in pos else (-1 if "B" in pos else 0)
        z = 1 if "U" in pos else (-1 if "D" in pos else 0)
        m[((x, y, z), pos[0])] = code
    return m


# --------------------------------------------------------------------------- #
# xlsx 解析
# --------------------------------------------------------------------------- #
@dataclass
class RawRow:
    type: str
    sheet: str
    master: bool
    code: str
    alg: str
    comm: str
    hand: str
    note: str
    alt: str
    alt_comm: str
    alt2: str
    alt2_comm: str


def _cell(row, i) -> str:
    if i is None or row is None or i >= len(row):
        return ""
    return norm_text(row[i])


def parse_workbook(path: str, type_: str, report: dict) -> List[RawRow]:
    wb = openpyxl.load_workbook(path, data_only=True)
    out: List[RawRow] = []
    for ws in wb.worksheets:
        title = ws.title
        if title in SKIP_SHEETS:
            continue
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            report["sheets"][title] = 0
            continue
        hdr = [norm_text(c) for c in rows[0]]
        i_code = hdr.index("编码") if "编码" in hdr else None
        i_alg = hdr.index("公式") if "公式" in hdr else None
        if i_code is None or i_alg is None:
            report["anomalies"].append(f"{title}: 缺 编码/公式 表头，跳过")
            report["sheets"][title] = 0
            continue
        i_comm = hdr.index("交换子") if "交换子" in hdr else None
        i_hand = hdr.index("起手") if "起手" in hdr else None
        i_note = hdr.index("备注") if "备注" in hdr else None
        i_alt = hdr.index("备选") if "备选" in hdr else None
        # 「备选」之后没有表头的三列依次是：备选交换子 / 备选2 / 备选2交换子
        i_altc = i_alt + 1 if i_alt is not None else None
        i_alt2 = i_alt + 2 if i_alt is not None else None
        i_alt2c = i_alt + 3 if i_alt is not None else None

        n = 0
        blank = 0
        for row in rows[1:]:
            if row is None:
                blank += 1
                continue
            code = _cell(row, i_code)
            alg = _cell(row, i_alg)
            if not code and not alg:
                blank += 1
                continue
            if not code or not alg:
                report["anomalies"].append(
                    f"{title}: 编码/公式 不完整的行 code={code!r} alg={alg!r}")
                if not alg:
                    continue
            if len(code) != 2:
                report["anomalies"].append(f"{title}: 编码长度异常 {code!r}")
                continue
            n += 1
            out.append(RawRow(
                type=type_, sheet=title, master=title in MASTER_SHEETS,
                code=code.upper(), alg=alg,
                comm=_cell(row, i_comm), hand=_cell(row, i_hand),
                note=_cell(row, i_note),
                alt=_cell(row, i_alt), alt_comm=_cell(row, i_altc),
                alt2=_cell(row, i_alt2), alt2_comm=_cell(row, i_alt2c),
            ))
        report["sheets"][title] = n
        report.setdefault("blank_rows", {})[title] = blank
    return out


def parse_letter_sheet(path: str) -> List[str]:
    """解析「编码」页的字母图，返回图中出现的 A–Z 单字母（面名 上/下/前/… 忽略）。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    if "编码" not in wb.sheetnames:
        return []
    ws = wb["编码"]
    letters = []
    for row in ws.iter_rows(values_only=True):
        for c in row:
            t = norm_text(c)
            if len(t) == 1 and "A" <= t.upper() <= "Z":
                letters.append(t.upper())
    return letters


def piece_layer(pos: str) -> str:
    """贴纸所属块所在层：U 层 / D 层 / 中层(E)。"""
    if "U" in pos:
        return "U"
    if "D" in pos:
        return "D"
    return "E"


# --------------------------------------------------------------------------- #
# 合并
# --------------------------------------------------------------------------- #
@dataclass
class Case:
    type: str
    code: str
    cats: List[str] = field(default_factory=list)
    alg: str = ""
    comm: str = ""
    hand: str = ""
    note: str = ""
    alt: str = ""
    alt_comm: str = ""
    alt2: str = ""
    alt2_comm: str = ""
    conflicts: List[str] = field(default_factory=list)


def merge(rows: List[RawRow], type_: str, report: dict) -> "OrderedDict[str, Case]":
    cases: "OrderedDict[str, Case]" = OrderedDict()
    for r in rows:
        c = cases.get(r.code)
        if c is None:
            c = Case(type=type_, code=r.code)
            cases[r.code] = c
        if not r.master:
            cat = CAT_ALIAS.get(r.sheet, r.sheet)
            if cat not in c.cats:
                c.cats.append(cat)
        for name in ("alg", "comm", "hand", "note", "alt", "alt_comm", "alt2", "alt2_comm"):
            new = getattr(r, name)
            if not new:
                continue
            if name in ("comm", "alt_comm", "alt2_comm") and is_placeholder(new):
                report["placeholder_" + name] += 1
                continue
            old = getattr(c, name)
            if not old:
                setattr(c, name, new)
            elif old != new:
                c.conflicts.append(f"{name}: {old!r} != {new!r}")
    for c in cases.values():
        if c.alt2 and not c.alt:
            c.alt, c.alt_comm = c.alt2, c.alt2_comm
            report["alt2_promoted"] += 1
        elif c.alt2:
            report["alt2_dropped"] += 1
        if not c.cats:
            c.cats = [FALLBACK_CAT]
            report["fallback_cat"] += 1
        order = CAT_ORDER[type_]
        c.cats.sort(key=lambda x: order.index(x) if x in order else len(order))
        if c.conflicts:
            report["conflicts"].extend(f"{c.code}: {x}" for x in c.conflicts)
    return cases


# --------------------------------------------------------------------------- #
# 映射推导
# --------------------------------------------------------------------------- #
def vote_mapping(cases, app_index, type_, report) -> Dict[str, Counter]:
    """公式投票：浅梦 2 字母码 → App 键，按位置投票。"""
    votes: Dict[str, Counter] = defaultdict(Counter)
    hit_codes = set()
    for c in cases.values():
        key = None
        for a in (c.alg, c.alt, c.alt2):
            if not a:
                continue
            k = app_index.get((type_, loose(a)))
            if k:
                key = k
                break
        if not key:
            continue
        hit_codes.add(c.code)
        votes[c.code[0]][key[1]] += 1
        votes[c.code[1]][key[2]] += 1
    report[f"vote_hit_{type_}"] = len(hit_codes)
    return votes


def majority(votes: Dict[str, Counter], report, type_) -> Optional[Dict[str, str]]:
    out, bad = {}, []
    for L in sorted(votes):
        top = votes[L].most_common()
        if len(top) > 1 and top[0][1] == top[1][1]:
            bad.append(f"{L}: 平票 {top[:3]}")
            continue
        out[L] = top[0][0]
    if bad:
        report[f"vote_tie_{type_}"] = bad
    return out or None


def is_bijection(m: Dict[str, str], alphabet: str) -> bool:
    if set(m) != set(alphabet):
        return False
    return len(set(m.values())) == len(m) and set(m.values()) == set(alphabet)


def finalize(mapping: Dict[str, str], alphabet: str, buffer_src: str, buffer_dst: str,
             forced: Dict[str, str]) -> Dict[str, str]:
    """补齐未定字母（恒等），并强制 缓冲位映射 + 指定字母对。"""
    f = dict(mapping)
    f[buffer_src] = buffer_dst
    for k, v in forced.items():
        f[k] = v
    for L in alphabet:
        f.setdefault(L, L)
    # 若出现非双射，按 forced → 恒等 的顺序修复
    for _ in range(4):
        used = Counter(f.values())
        dup = [k for k, v in used.items() if v > 1]
        if not dup:
            break
        for L in list(f):
            if L in forced or L == buffer_src:
                continue
            if used[f[L]] > 1:
                f[L] = L
    return f


def usable_vote(mapping, alphabet, buffer_src, buffer_dst, forced) -> bool:
    """投票结果 + 缓冲约束 + 恒等补齐后是否为合法双射。"""
    if not mapping:
        return False
    f = dict(mapping)
    f[buffer_src] = buffer_dst
    f.update(forced)
    for L in alphabet:
        f.setdefault(L, L)
    return is_bijection(f, alphabet)


# --------------------------------------------------------------------------- #
# 主流程
# --------------------------------------------------------------------------- #
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="浅梦公式集 → alg818_qm.txt")
    ap.add_argument("--check", action="store_true", help="只自检，不写文件")
    ap.add_argument("--quiet", action="store_true", help="只打印结论")
    args = ap.parse_args(argv)

    report = {"sheets": {}, "anomalies": [], "conflicts": [], "fallback_cat": 0,
              "alt2_dropped": 0, "alt2_promoted": 0,
              "placeholder_comm": 0, "placeholder_alt_comm": 0, "placeholder_alt2_comm": 0}
    say = (lambda *a: None) if args.quiet else print

    # ---- 输入存在性 ----
    for p in (XLSX_EDGE, XLSX_CORNER, ALG818,
              os.path.join(JSOND, "cornerPosToCode.json"),
              os.path.join(JSOND, "edgePosToCode.json")):
        if not os.path.exists(p):
            print(f"[FATAL] 缺少输入：{p}", file=sys.stderr)
            return 2

    corner_pos = json.load(open(os.path.join(JSOND, "cornerPosToCode.json"), encoding="utf-8"))
    edge_pos = json.load(open(os.path.join(JSOND, "edgePosToCode.json"), encoding="utf-8"))
    letter_slots = {"C": build_letter_slots(corner_pos), "E": build_letter_slots(edge_pos)}
    pos_of = {"C": {v: k for k, v in corner_pos.items()},
              "E": {v: k for k, v in edge_pos.items()}}
    alphabet = {"C": "".join(sorted(corner_pos.values())),
                "E": "".join(sorted(edge_pos.values()))}
    buffer_src = {"C": "J", "E": "A"}   # 浅梦：角 J(UFR) / 棱 A(UF)
    buffer_dst = {"C": "J", "E": "G"}   # App ：角 J(UFR) / 棱 G(UR)

    # ---- alg818.txt ----
    app_rows = []
    with open(ALG818, encoding="utf-8") as fh:
        for line in fh:
            line = line.rstrip("\n")
            if not line:
                continue
            f = line.split("\t")
            if len(f) < 4:
                report["anomalies"].append(f"alg818: 字段不足 {line[:40]!r}")
                continue
            app_rows.append((f[0], f[1], f[2], f[3:]))
    app_keys = {"C": set(), "E": set()}
    app_index = {}
    for t, key, rec, algs in app_rows:
        app_keys[t].add(key)
        for a in algs:
            if a:
                app_index.setdefault((t, loose(a)), key)
    say(f"alg818.txt: {len(app_rows)} 行, C={len(app_keys['C'])} E={len(app_keys['E'])}, "
        f"算法索引 {len(app_index)} 条")

    # ---- xlsx ----
    edge_rows = parse_workbook(XLSX_EDGE, "E", report)
    corner_rows = parse_workbook(XLSX_CORNER, "C", report)
    edge_cases = merge(edge_rows, "E", report)
    corner_cases = merge(corner_rows, "C", report)
    cases = {"C": corner_cases, "E": edge_cases}

    say("\n-- 各 sheet 解析行数 --")
    for t, rows in (("E", edge_rows), ("C", corner_rows)):
        per = Counter(r.sheet for r in rows)
        blanks = report.get("blank_rows", {})
        say(f"  {t}: " + "  ".join(
            f"{k}={v}" + (f"(空行{blanks[k]})" if blanks.get(k) else "")
            for k, v in per.items()) + f"  | 有效合计 {len(rows)}")
    say(f"  去重后：C={len(corner_cases)}  E={len(edge_cases)}  （期望 378 / 440）")
    for t in "CE":
        multi = sum(1 for c in cases[t].values() if len(c.cats) > 1)
        say(f"  分类并集：{t} 中跨多个分类页的 case {multi} 条")

    # ---- 编码页字母图 ----
    for t, path in (("E", XLSX_EDGE), ("C", XLSX_CORNER)):
        letters = parse_letter_sheet(path)
        cnt = Counter(letters)
        ok = (sorted(cnt) == sorted(alphabet[t])) and all(v == 1 for v in cnt.values())
        say(f"  编码页[{t}]: {len(letters)} 个 A–Z 字母, 唯一 {len(cnt)}, "
            f"每字母恰好一次={all(v == 1 for v in cnt.values())}, "
            f"与 posToCode 字母集一致={ok}")
        if not ok:
            report["anomalies"].append(f"编码页[{t}] 字母集不一致: {sorted(cnt)}")

    # ---- 公式投票 ----
    votes = {t: vote_mapping(cases[t], app_index, t, report) for t in "CE"}
    maj = {t: majority(votes[t], report, t) for t in "CE"}
    for t in "CE":
        say(f"\n-- 公式投票 [{t}] -- 命中 case {report[f'vote_hit_{t}']}/{len(cases[t])}")
        if maj[t]:
            agree = sum(1 for k, v in maj[t].items() if k == v)
            say(f"   多数票 {len(maj[t])} 个字母，其中恒等 {agree} 个："
                + " ".join(f"{k}→{v}" for k, v in sorted(maj[t].items())))
            if not is_bijection(dict(maj[t], **{L: L for L in alphabet[t] if L not in maj[t]}),
                                alphabet[t]):
                say("   → 多数票不是双射（同一条公式在 App 里对应的是把 3-cycle 从 UR "
                    "重新起头的键，字母顺序被旋转），不能直接采用。")
        else:
            say("   无可用多数票。")

    # ---- 最终映射 ----
    # 角块：两边缓冲同为 J(UFR)，投票给出恒等，J/K/L 从不出现在编码里，恒等补齐。
    forced_c = {}
    ok_c = usable_vote(maj["C"], alphabet["C"], buffer_src["C"], buffer_dst["C"], forced_c)
    f_map = {"C": finalize(maj["C"] if ok_c else {}, alphabet["C"],
                           buffer_src["C"], buffer_dst["C"], forced_c)}
    say(f"   角块投票可用={ok_c}")
    # 棱块：浅梦缓冲 A(UF) → App 缓冲 G(UR)。App 键里不能出现 G/H，而浅梦编码里
    # G/H 是普通字母，因此必须把 A/B 与 G/H 对调（其余 20 字母恒等，学员的词对不变）。
    forced_e = {"A": "G", "B": "H", "G": "A", "H": "B"}
    ok_e = usable_vote(maj["E"], alphabet["E"], buffer_src["E"], buffer_dst["E"], forced_e)
    f_map["E"] = finalize(maj["E"] if ok_e else {}, alphabet["E"],
                          buffer_src["E"], buffer_dst["E"], forced_e)
    say(f"   棱块投票可用={ok_e}；最终采用 "
        f"{'投票+约束' if ok_e else '缓冲对应 + 双射约束（A/B ↔ G/H 对调，其余恒等）'}")

    for t in "CE":
        if not is_bijection(f_map[t], alphabet[t]):
            print(f"[FATAL] 映射 {t} 不是双射：{f_map[t]}", file=sys.stderr)
            return 3
    say("\n-- 最终映射（浅梦 → App）--")
    for t in "CE":
        say(f"  {t}: " + " ".join(f"{k}→{v}" for k, v in sorted(f_map[t].items())))

    # ---- 生成 KEY ----
    keyed = {}
    for t in "CE":
        for c in cases[t].values():
            unknown = [ch for ch in c.code if ch not in f_map[t]]
            if unknown:
                print(f"[FATAL] {t} {c.code}: 编码含未知字母 {unknown}", file=sys.stderr)
                return 5
            k = buffer_dst[t] + f_map[t][c.code[0]] + f_map[t][c.code[1]]
            keyed[(t, c.code)] = k

    # ---- 校验 1：键集合 ----
    problems = []
    for t in "CE":
        got = Counter(keyed[(t, c.code)] for c in cases[t].values())
        dup = [k for k, v in got.items() if v > 1]
        want = app_keys[t]
        missing = sorted(want - set(got))
        extra = sorted(set(got) - want)
        say(f"\n-- 校验1 键集合 [{t}] -- 生成 {len(got)} 个键, 重复 {len(dup)}, "
            f"缺 {len(missing)}, 多 {len(extra)}")
        if dup:
            say(f"   重复键: {dup[:8]}")
        if missing:
            say(f"   缺失: {missing[:8]}")
        if extra:
            say(f"   多余: {extra[:8]}")
        if dup or missing or extra:
            problems.append(f"{t} 键集合不匹配")

    # ---- 校验 2：魔方模拟（编码 ↔ 公式 一致）----
    sim_bad, sim_ok, tok_bad = [], 0, Counter()
    for t in "CE":
        L = letter_slots[t]
        buf = buffer_src[t]
        for c in cases[t].values():
            algs = [a for a in (c.alg,) if a]
            if not algs:
                continue
            try:
                perm = alg_perm(norm_text(c.alg))
            except Exception as exc:  # noqa: BLE001
                sim_bad.append((t, c.code, "解析失败", str(exc)))
                continue
            for tok in norm_text(c.alg).split():
                if parse_token(tok) is None:
                    tok_bad[tok] += 1
            names = [[L[s] for s in cyc] for cyc in cycles_of(perm)]
            hit = [n for n in names if buf in n]
            if not hit:
                sim_bad.append((t, c.code, "无缓冲位循环", names))
                continue
            cyc = hit[0]
            i = cyc.index(buf)
            rot = "".join(cyc[i:] + cyc[:i])
            if rot == buf + c.code:
                sim_ok += 1
            else:
                sim_bad.append((t, c.code, "循环不符", rot))
    say(f"\n-- 校验2 魔方模拟 -- 一致 {sim_ok}/{sum(len(cases[t]) for t in 'CE')}, "
        f"异常 {len(sim_bad)}")
    for b in sim_bad[:10]:
        say(f"   {b}")
    if tok_bad:
        say(f"   未知转动记号: {dict(tok_bad)}")
    if sim_bad:
        problems.append("模拟校验失败")

    # ---- 校验 3：语义抽检 ----
    say("\n-- 校验3 语义抽检（每类 ≥2 条）--")
    face_bad, piece_bad = [], []
    for t in "CE":
        for cat in CAT_ORDER[t]:
            group = sorted([c for c in cases[t].values() if cat in c.cats],
                           key=lambda c: c.code)
            if not group:
                continue
            sample = group[:2]
            if len(group) > 3:
                sample.append(group[-1])
            faces = Counter()
            pieces = Counter()
            for c in group:
                for ch in c.code:
                    faces[pos_of[t][ch][0]] += 1
                    pieces[piece_layer(pos_of[t][ch])] += 1
            for c in sample:
                k = keyed[(t, c.code)]
                px, py = pos_of[t][c.code[0]], pos_of[t][c.code[1]]
                mx, my = f_map[t][c.code[0]], f_map[t][c.code[1]]
                say(f"   {t}/{cat:<10} {c.code}→{k}  "
                    f"{c.code[0]}@{px}[{px[0]}面/{piece_layer(px)}块]→{mx}@{pos_of[t][mx]}"
                    f"[{pos_of[t][mx][0]}面/{piece_layer(pos_of[t][mx])}块]  "
                    f"{c.code[1]}@{py}[{py[0]}面/{piece_layer(py)}块]→{my}@{pos_of[t][my]}"
                    f"[{pos_of[t][my][0]}面/{piece_layer(pos_of[t][my])}块]")
            say(f"      该类 {len(group)} 条；编码字母贴纸面分布 {dict(faces)}，"
                f"所属块层分布 {dict(pieces)}；说明：{CAT_HINT.get(cat, '')}")
    for t in "CE":
        for L in alphabet[t]:
            src, dst = L, f_map[t][L]
            if pos_of[t][dst][0] != pos_of[t][src][0]:
                face_bad.append(f"{L}@{pos_of[t][src]}({pos_of[t][src][0]}面)→"
                                f"{dst}@{pos_of[t][dst]}({pos_of[t][dst][0]}面)")
            if piece_layer(pos_of[t][dst]) != piece_layer(pos_of[t][src]):
                piece_bad.append(f"{L}→{dst}")
    say(f"   贴纸所在面保持不变: {24 * 2 - len(face_bad)}/{24 * 2} 字母·类；变化: {face_bad}")
    say(f"   贴纸所属块所在层保持不变: {24 * 2 - len(piece_bad)}/{24 * 2} 字母·类；变化: {piece_bad}")
    if piece_bad:
        problems.append("块层语义未保持")
    if face_bad:
        say("   注：B/H 互换会让 H(R 面)→B(F 面)，但 B 从不出现在浅梦编码里，"
            "实际受影响只有 H 一个字母；A↔G 同为 U 面，不受影响。")

    # ---- 校验 4：与 alg818.txt 同键算法重合度（信息性）----
    same_key = {"C": 0, "E": 0}
    any_key = {"C": 0, "E": 0}
    for t in "CE":
        for c in cases[t].values():
            k = keyed[(t, c.code)]
            got = app_index.get((t, loose(c.alg)))
            if got is not None:
                any_key[t] += 1
                if got == k:
                    same_key[t] += 1
    say(f"\n-- 校验4 信息性 -- 主公式在 alg818.txt 中能查到: "
        f"C={any_key['C']}/{len(cases['C'])}  E={any_key['E']}/{len(cases['E'])}；"
        f"其中键也相同: C={same_key['C']}  E={same_key['E']}")
    say("   角块两边缓冲同为 UFR(J)，case 完全同构；棱块浅梦缓冲 UF(A)、App 缓冲 UR(G)，"
        "只有 3-cycle 穿过 UR 的 case 会重合，其余为形式重标号。")

    # ---- 写文件 ----
    lines = []
    key_re = {"C": re.compile(r"^J[A-Z]{2}$"), "E": re.compile(r"^G[A-Z]{2}$")}
    for t in ("C", "E"):
        for c in sorted(cases[t].values(), key=lambda x: keyed[(x.type, x.code)]):
            k = keyed[(t, c.code)]
            if not key_re[t].match(k):
                print(f"[FATAL] KEY 不合法: {k!r} ({t} {c.code})", file=sys.stderr)
                return 4
            cat = "|".join(c.cats)
            fields = [t, k, cat, c.hand, c.comm, c.note, c.alg, c.alt, c.alt_comm]
            for i, v in enumerate(fields):
                if "\t" in v or "\n" in v or "\r" in v:
                    print(f"[FATAL] 字段含控制字符: {c.code} idx={i}", file=sys.stderr)
                    return 4
            lines.append("\t".join(fields))
    data = "\n".join(lines) + "\n"
    blob = data.encode("utf-8")

    field_dist = Counter(len(l.split("\t")) for l in lines)
    say(f"\n-- 输出 -- {len(lines)} 行, {len(blob)} 字节, 字段数分布 {dict(field_dist)}, "
        f"sha256={hashlib.sha256(blob).hexdigest()[:16]}")
    if len(lines) != 818:
        problems.append(f"行数 {len(lines)} != 818")
    if set(field_dist) != {9}:
        problems.append(f"字段数分布异常 {dict(field_dist)}")

    if not args.check:
        old = None
        if os.path.exists(OUTFILE):
            old = open(OUTFILE, "rb").read()
        if old == blob:
            say("   文件内容未变（幂等）")
        else:
            tmp = OUTFILE + ".tmp"
            with open(tmp, "wb") as fh:
                fh.write(blob)
            os.replace(tmp, OUTFILE)
            say(f"   已写入 {OUTFILE}")

    # ---- 汇总 ----
    say("\n-- 异常汇总 --")
    for a in report["anomalies"]:
        say(f"   ! {a}")
    for c in report["conflicts"][:20]:
        say(f"   ! 冲突 {c}")
    say(f"   无分类兜底 {report['fallback_cat']} 条；总表/分类页字段冲突 {len(report['conflicts'])} 条")
    say(f"   第二备选被丢弃 {report['alt2_dropped']} 条（格式只有一组 ALT），"
        f"无第一备选而提升第二备选 {report['alt2_promoted']} 条")
    say(f"   'Not found.' 占位转空串：交换子 {report['placeholder_comm']} 处，"
        f"备选交换子 {report['placeholder_alt_comm']} 处，"
        f"第二备选交换子 {report['placeholder_alt2_comm']} 处")
    say(f"   空行跳过：" + "  ".join(f"{k}={v}" for k, v in report.get("blank_rows", {}).items() if v))
    say(f"\n{'PASS' if not problems else 'FAIL'}: " + ("；".join(problems) if problems else
        "818/818 全覆盖，键集合与 alg818.txt 完全相等"))
    return 0 if not problems else 1


if __name__ == "__main__":
    sys.exit(main())
