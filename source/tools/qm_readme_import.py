# -*- coding: utf-8 -*-
"""从「浅梦」角块/棱块分类公式集 xlsx 提取说明文本，生成 App 内可渲染的纯文本资产。

输入（只读）：
  C:\\Users\\Estiavl\\Desktop\\RUBIK\\浅梦's 角块UFR分类公式集 202602版.xlsx
  C:\\Users\\Estiavl\\Desktop\\RUBIK\\浅梦's 棱块UF分类公式集 202602版.xlsx
输出：
  D:\\Program\\dsh\\rebuild\\src\\assets\\qm_readme.txt

行协议（App 端按此解析）：
  #V1        固定首行
  [S] xxx    一级标题
  [G] xxx    分组标题（分页名）
  [P] xxx    段落
  [I] xxx    列表条目（App 加圆点）
  [L] xxx    链接

用法：
  python qm_readme_import.py            # 生成资产
  python qm_readme_import.py --check    # 只校验现有资产与重新生成内容是否逐字节一致
  python qm_readme_import.py --quiet    # 不打印统计

幂等：同一份输入重复运行，输出逐字节一致（UTF-8 无 BOM / LF / 末尾单个换行）。
"""

from __future__ import annotations

import argparse
import hashlib
import os
import re
import sys

import openpyxl

# ---------------------------------------------------------------- 路径常量

XLSX_CORNER = (
    "C:\\Users\\Estiavl\\Desktop\\RUBIK\\"
    "浅梦's 角块UFR分类公式集 202602版.xlsx"
)
XLSX_EDGE = (
    "C:\\Users\\Estiavl\\Desktop\\RUBIK\\"
    "浅梦's 棱块UF分类公式集 202602版.xlsx"
)
OUT_PATH = "D:\\Program\\dsh\\rebuild\\src\\assets\\qm_readme.txt"

TAG_CORNER = "角块（UFR）"
TAG_EDGE = "棱块（UF）"

HEADER = "#V1"
LINE_RE = re.compile(r"^(#V1|\[S\] |\[G\] |\[P\] |\[I\] |\[L\] )")

SKIP_SHEETS = {"说明", "编码"}
SUMMARY_SUFFIX = "总表"

PAGE_EMPTY = "（原表这一页没有整页说明）"

# 说明 sheet A 列小标题 -> 内部 key（归一化后比较）
LABEL_MAP = {
    "简介": "intro",
    "在线文档链接": "link",
    "编码部分": "code",
    "公式部分": "formula",
    "备注部分": "note",
    "起手部分": "start",
    "分页部分": "pages",  # 原表无正文，跳过；分页说明由各分页 sheet 备注生成
}

# 输出顺序（分页说明单独处理）
SECTION_ORDER = [
    ("intro", "简介"),
    ("link", "在线文档"),
    ("code", "编码部分"),
    ("formula", "公式部分"),
    ("note", "备注部分"),
    ("start", "起手部分"),
]

_QUOTE_CHARS = "“”\"'‘’"


# ---------------------------------------------------------------- 文本工具


def _norm_newlines(value) -> str:
    """统一换行符并去掉首尾空白（含换行）。"""
    s = "" if value is None else str(value)
    s = s.replace("\r\n", "\n").replace("\r", "\n").replace("\u2028", "\n")
    s = s.replace("\u00a0", " ")
    return s.strip()


def split_paragraphs(value) -> list:
    """按换行拆段：每行一段，空行丢弃，行内首尾空白去掉。"""
    s = _norm_newlines(value)
    if not s:
        return []
    out = []
    for line in s.split("\n"):
        line = line.strip()
        if line:
            out.append(line)
    return out


def split_list_items(value) -> list:
    """起手部分：按 '；' 与换行拆成条目，去掉首尾空白与句末分号。"""
    s = _norm_newlines(value)
    if not s:
        return []
    out = []
    for chunk in re.split(r"[；;\n]", s):
        item = chunk.strip().strip("；;").strip()
        if item:
            out.append(item)
    return out


def norm_label(value) -> str:
    """归一化说明 sheet 的 A 列小标题：去引号/空白/结尾冒号。"""
    s = "" if value is None else str(value)
    for ch in _QUOTE_CHARS:
        s = s.replace(ch, "")
    s = s.replace(" ", "").replace("\u3000", "").replace("\t", "")
    s = s.strip()
    s = s.rstrip("：:").strip()
    return s


def assert_single_line(text: str, where: str) -> None:
    if "\n" in text or "\r" in text:
        raise ValueError("%s 含有换行符，无法写入单行协议：%r" % (where, text))


# ---------------------------------------------------------------- 读取 workbook


def read_readme_sheet(ws) -> dict:
    """读「说明」sheet，返回 {key: 原始正文}。"""
    found = {}
    for row in ws.iter_rows(min_row=1, max_col=2, values_only=True):
        label_raw = row[0] if len(row) > 0 else None
        body = row[1] if len(row) > 1 else None
        label = norm_label(label_raw)
        if not label:
            continue
        key = LABEL_MAP.get(label)
        if key is None:
            raise ValueError("未知的说明小标题：%r（归一化为 %r）" % (label_raw, label))
        if key == "pages":
            # 原表没有正文，跳过（分页说明由各分页 sheet 的「备注」列生成）
            continue
        if key in found:
            raise ValueError("说明 sheet 中小标题重复：%r" % (label_raw,))
        found[key] = body
    return found


def read_pages(wb) -> list:
    """按 workbook 实际 sheet 顺序提取各分页的整页备注。

    返回 [(sheet 名, [段落...]), ...]；跳过 说明 / 编码 / 总表。
    """
    pages = []
    for ws in wb.worksheets:
        title = str(ws.title)
        stripped = title.strip()
        if stripped in SKIP_SHEETS or stripped.endswith(SUMMARY_SUFFIX):
            continue

        # 找表头「备注」列（表头可能带尾空格）
        note_col = None
        for col in range(1, ws.max_column + 1):
            head = ws.cell(row=1, column=col).value
            if head is not None and str(head).strip() == "备注":
                note_col = col
                break
        if note_col is None:
            pages.append((title, []))
            continue

        # 取第一个非空的备注单元格
        paras = []
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=note_col).value
            if cell is None:
                continue
            if not str(cell).strip():
                continue
            paras = split_paragraphs(cell)
            break
        pages.append((title, paras))
    return pages


def load_workbook_data(path: str, tag: str) -> dict:
    if not os.path.isfile(path):
        raise FileNotFoundError("找不到输入文件：%s" % path)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    try:
        if "说明" not in wb.sheetnames:
            raise ValueError("%s 中找不到「说明」sheet" % path)
        raw = read_readme_sheet(wb["说明"])
        pages = read_pages(wb)
    finally:
        wb.close()

    data = {
        "tag": tag,
        "path": path,
        "intro": split_paragraphs(raw.get("intro")),
        "link": _norm_newlines(raw.get("link")),
        "code": split_paragraphs(raw.get("code")),
        "formula": split_paragraphs(raw.get("formula")),
        "note": split_paragraphs(raw.get("note")),
        "start": split_list_items(raw.get("start")),
        "pages": pages,
    }
    return data


# ---------------------------------------------------------------- 组装资产


def emit_para_section(lines: list, title: str, key: str, a: dict, b: dict, kind: str) -> None:
    """段落型 section：两份相同时只写一份，不同时按 角块/棱块 分块。"""
    pa, pb = a[key], b[key]
    lines.append("[S] " + title)
    if pa == pb:
        lines.extend("[%s] %s" % (kind, p) for p in pa)
        return
    lines.append("[P] " + a["tag"])
    lines.extend("[%s] %s" % (kind, p) for p in pa)
    lines.append("[P] " + b["tag"])
    lines.extend("[%s] %s" % (kind, p) for p in pb)


def emit_link_section(lines: list, title: str, a: dict, b: dict) -> None:
    lines.append("[S] " + title)
    la, lb = a["link"], b["link"]
    if la == lb:
        if la:
            lines.append("[L] " + la)
        return
    for tag, link in ((a["tag"], la), (b["tag"], lb)):
        if not link:
            continue
        lines.append("[I] " + tag)
        lines.append("[L] " + link)


def emit_page_section(lines: list, title: str, pages: list) -> None:
    lines.append("[S] " + title)
    for name, paras in pages:
        lines.append("[G] " + name)
        if paras:
            lines.extend("[P] " + p for p in paras)
        else:
            lines.append("[P] " + PAGE_EMPTY)


def build_text() -> str:
    corner = load_workbook_data(XLSX_CORNER, TAG_CORNER)
    edge = load_workbook_data(XLSX_EDGE, TAG_EDGE)

    lines = [HEADER]
    for key, title in SECTION_ORDER:
        if key == "link":
            emit_link_section(lines, title, corner, edge)
        elif key == "start":
            emit_para_section(lines, title, key, corner, edge, "I")
        else:
            emit_para_section(lines, title, key, corner, edge, "P")

    emit_page_section(lines, "角块分页说明", corner["pages"])
    emit_page_section(lines, "棱块分页说明", edge["pages"])

    # 每行都不允许换行符
    for i, line in enumerate(lines):
        assert_single_line(line, "第 %d 行" % (i + 1))

    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------- 统计 / 校验


def collect_stats(text: str) -> dict:
    sections = []
    g_counts = {}
    cur = None
    start_items = 0
    in_start = False
    for line in text.split("\n"):
        if line.startswith("[S] "):
            cur = line[4:]
            sections.append(cur)
            g_counts[cur] = 0
            in_start = cur == "起手部分"
        elif line.startswith("[G] ") and cur is not None:
            g_counts[cur] += 1
        elif line.startswith("[I] ") and in_start:
            start_items += 1
    data = text.encode("utf-8")
    return {
        "sections": sections,
        "g_counts": g_counts,
        "start_items": start_items,
        "bytes": len(data),
        "sha256": hashlib.sha256(data).hexdigest(),
        "lines": len(text.split("\n")) - 1,
    }


def print_stats(text: str, path: str) -> None:
    st = collect_stats(text)
    print("[stats] 输出: %s" % path)
    print("[stats] [S] 一级标题: %d 个" % len(st["sections"]))
    for name in st["sections"]:
        print("        [S] %s" % name)
    print("[stats] [G] 分组: %d 个" % sum(st["g_counts"].values()))
    for name, n in st["g_counts"].items():
        if n:
            print("        %s: %d" % (name, n))
    print("[stats] 起手部分 [I] 条目: %d 条" % st["start_items"])
    print("[stats] 行数: %d" % st["lines"])
    print("[stats] 字节数: %d" % st["bytes"])
    print("[stats] sha256: %s" % st["sha256"])
    print("[stats] 分页 [G] 名称:")
    cur = None
    for line in text.split("\n"):
        if line.startswith("[S] "):
            cur = line[4:]
        elif line.startswith("[G] "):
            print("        %s / %s" % (cur, line[4:]))


def diff_summary(old: bytes, new: bytes) -> str:
    """打印逐字节不一致的摘要。"""
    old_text = old.decode("utf-8", errors="replace")
    new_text = new.decode("utf-8", errors="replace")
    old_lines = old_text.split("\n")
    new_lines = new_text.split("\n")
    out = []
    out.append("  现有: %d 字节, %d 行, sha256=%s"
               % (len(old), len(old_lines), hashlib.sha256(old).hexdigest()))
    out.append("  新生成: %d 字节, %d 行, sha256=%s"
               % (len(new), len(new_lines), hashlib.sha256(new).hexdigest()))
    if old.startswith(b"\xef\xbb\xbf"):
        out.append("  差异: 现有文件带 UTF-8 BOM")
    if b"\r" in old:
        out.append("  差异: 现有文件含 CR 字符（非纯 LF）")
    diffs = 0
    for i in range(max(len(old_lines), len(new_lines))):
        a = old_lines[i] if i < len(old_lines) else "<缺行>"
        b = new_lines[i] if i < len(new_lines) else "<缺行>"
        if a != b:
            diffs += 1
            if diffs <= 10:
                out.append("  行 %d:\n    现有: %r\n    新生成: %r" % (i + 1, a, b))
    out.append("  不一致行数: %d" % diffs)
    return "\n".join(out)


# ---------------------------------------------------------------- 主流程


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="提取两份公式集的说明，生成 qm_readme.txt 资产")
    ap.add_argument("--check", action="store_true",
                    help="只校验现有资产与重新生成内容是否逐字节一致（不一致退出码 1）")
    ap.add_argument("--quiet", action="store_true", help="不打印统计")
    ap.add_argument("--out", default=OUT_PATH, help="输出路径（默认 %s）" % OUT_PATH)
    args = ap.parse_args(argv)

    text = build_text()
    new_bytes = text.encode("utf-8")

    if args.check:
        if not os.path.isfile(args.out):
            print("[check] 资产不存在: %s" % args.out)
            return 1
        with open(args.out, "rb") as fh:
            old_bytes = fh.read()
        if old_bytes == new_bytes:
            if not args.quiet:
                st = collect_stats(text)
                print("[check] OK 逐字节一致 (%d 字节, sha256=%s)"
                      % (st["bytes"], st["sha256"]))
            return 0
        print("[check] FAIL 资产与重新生成内容不一致: %s" % args.out)
        print(diff_summary(old_bytes, new_bytes))
        return 1

    out_dir = os.path.dirname(args.out)
    if out_dir and not os.path.isdir(out_dir):
        os.makedirs(out_dir)
    with open(args.out, "wb") as fh:
        fh.write(new_bytes)

    if not args.quiet:
        print_stats(text, args.out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
